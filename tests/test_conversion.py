"""End-to-end conversion checks, including an independent OLE/record reader."""
import json
import struct
import sys
import tempfile
import unittest
import zlib
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / "vendor/python")]

import olefile
from PIL import Image
from docx import Document
from docx.shared import Inches
from convert import convert, inspect_hwp


def records(data):
    result = []
    offset = 0
    while offset < len(data):
        header, = struct.unpack_from("<I", data, offset)
        offset += 4
        tag, level, size = header & 1023, (header >> 10) & 1023, header >> 20
        if size == 4095:
            size, = struct.unpack_from("<I", data, offset)
            offset += 4
        assert offset + size <= len(data), "Record extends outside stream"
        result.append((tag, level, data[offset:offset + size]))
        offset += size
    return result


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def md(self, text, name="input.md"):
        path = self.root / name
        path.write_text(text, encoding="utf-8")
        return path

    def verify_container(self, path):
        with olefile.OleFileIO(path) as ole:
            self.assertTrue(ole.exists("FileHeader"))
            self.assertTrue(ole.exists("DocInfo"))
            self.assertTrue(ole.exists("BodyText/Section0"))
            header = ole.openstream("FileHeader").read()
            self.assertTrue(header.startswith(b"HWP Document File"))
            self.assertEqual(header[35], 5)
            compressed = bool(struct.unpack_from("<I", header, 36)[0] & 1)
            section = ole.openstream("BodyText/Section0").read()
            docinfo = ole.openstream("DocInfo").read()
            if compressed:
                section = zlib.decompress(section, -15)
                docinfo = zlib.decompress(docinfo, -15)
            body = records(section)
            records(docinfo)
            self.assertTrue(any(tag == 66 for tag, _, _ in body))
            self.assertTrue(any(tag == 67 for tag, _, _ in body))
            return body

    def test_markdown_unicode_styles_table_and_order(self):
        source = self.md("# 제목 한글\n\n처음 **강조** *기울임* 😀\n\n"
                         "|열 하나|열 둘|\n|---|---|\n|값 가|값 나|\n\n마지막 문단\n")
        report = convert(source)
        self.assertEqual(report["tables"], 1)
        output = Path(report["output"])
        body = self.verify_container(output)
        text = inspect_hwp(output)["text"]
        for value in ["제목 한글", "강조", "기울임", "😀", "값 가", "값 나", "마지막 문단"]:
            self.assertIn(value, text)
        self.assertLess(text.index("처음"), text.index("값 가"))
        self.assertLess(text.index("값 나"), text.index("마지막 문단"))
        self.assertTrue(any(tag == 77 for tag, _, _ in body), "HWP table record missing")
        self.assertGreater(inspect_hwp(output)["char_shapes"], 7)

    def test_docx_image_table_merge_and_paragraph_order(self):
        picture = self.root / "image.png"
        Image.new("RGB", (120, 60), "navy").save(picture)
        doc = Document()
        doc.add_heading("DOCX 변환 확인", 1)
        doc.add_paragraph("본문 시작").add_run(" 굵은 텍스트").bold = True
        table = doc.add_table(rows=3, cols=3)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "가로 병합"
        table.cell(1, 0).merge(table.cell(2, 0)).text = "세로 병합"
        table.cell(2, 1).text = "아래 셀"
        doc.add_picture(str(picture), width=Inches(1.5))
        doc.add_paragraph("마지막 확인")
        source = self.root / "input.docx"
        doc.save(source)
        report = convert(source)
        output = Path(report["output"])
        self.assertEqual((report["tables"], report["images"]), (1, 1))
        self.verify_container(output)
        inspected = inspect_hwp(output)
        self.assertEqual(inspected["embedded_images"], 1)
        self.assertEqual(inspected["text"].count("세로 병합"), 1)
        self.assertIn("아래 셀", inspected["text"])
        self.assertIn("마지막 확인", inspected["text"])
        with olefile.OleFileIO(output) as ole:
            images = [entry for entry in ole.listdir() if entry[0] == "BinData"]
            self.assertEqual(len(images), 1)
            self.assertTrue(ole.openstream(images[0]).read().startswith(b"\x89PNG"))

    def test_markdown_image_embedded(self):
        Image.new("RGB", (40, 20), "red").save(self.root / "그림.png")
        report = convert(self.md("앞\n\n![그림](그림.png)\n\n뒤\n"))
        self.assertEqual(report["images"], 1)
        inspected = inspect_hwp(report["output"])
        self.assertEqual(inspected["embedded_images"], 1)
        self.assertIn("앞", inspected["text"])
        self.assertIn("뒤", inspected["text"])

    def test_existing_output_is_preserved_and_force_replaces(self):
        source = self.md("새 문서")
        output = source.with_suffix(".hwp")
        output.write_bytes(b"existing user document")
        with self.assertRaises(FileExistsError):
            convert(source)
        self.assertEqual(output.read_bytes(), b"existing user document")
        convert(source, force=True)
        self.verify_container(output)

    def test_strict_warnings_leave_no_output(self):
        source = self.md("![원격](https://example.invalid/image.png)")
        with self.assertRaises(ValueError):
            convert(source, strict=True)
        self.assertFalse(source.with_suffix(".hwp").exists())
        self.assertEqual(list(self.root.glob("hwp-convert-*")), [])

    def test_validation_and_empty_input(self):
        with self.assertRaises(ValueError):
            convert(self.root / "missing.md")
        with self.assertRaises(ValueError):
            convert(self.md("text"), self.root / "wrong.docx")
        result = convert(self.md("", "empty.md"))
        self.verify_container(result["output"])


if __name__ == "__main__":
    unittest.main()
