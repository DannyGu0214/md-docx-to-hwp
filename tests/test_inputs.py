"""Behavioral tests using temporary inputs, including generated OOXML fixtures."""
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "vendor" / "python"))

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from converter.inputs import parse_input


def paragraph_text(block):
    return "".join(run["text"] for run in block["runs"])


class InputTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.directory = Path(self.temp.name)
        self.assets = self.directory / "assets"

    def tearDown(self):
        self.temp.cleanup()

    def markdown(self, text):
        source = self.directory / "input.md"
        source.write_text(text, encoding="utf-8")
        return parse_input(source, self.assets)

    def save_docx(self, document):
        source = self.directory / "input.docx"
        document.save(source)
        return parse_input(source, self.assets)

    def image(self):
        source = self.directory / "그림.png"
        Image.new("RGB", (80, 40), (120, 30, 200)).save(source)
        return source

    def test_markdown_formats_lists_table_and_korean(self):
        parsed = self.markdown(
            "# 제목 한글\n\n문장 **굵게** *기울임* ~~삭제~~ `코드`\n\n"
            "3. 셋째\n4. 넷째\n\n- 항목\n  - 안쪽\n\n"
            "| 이름 | 값 |\n| :--- | ---: |\n| 대니 | 42 |\n")
        blocks = parsed["blocks"]
        self.assertEqual(blocks[0]["heading"], 1)
        self.assertEqual(paragraph_text(blocks[0]), "제목 한글")
        runs = blocks[1]["runs"]
        self.assertTrue(next(run for run in runs if run["text"] == "굵게")["bold"])
        self.assertTrue(next(run for run in runs if run["text"] == "기울임")["italic"])
        self.assertTrue(next(run for run in runs if run["text"] == "삭제")["strike"])
        self.assertEqual(next(run for run in runs if run["text"] == "코드")["font"], "Courier New")
        self.assertEqual([b["list_prefix"] for b in blocks if "list_prefix" in b],
                         ["3.", "4.", "•", "  •"])
        table = next(b for b in blocks if b["type"] == "table")
        self.assertEqual(len(table["rows"]), 2)
        self.assertEqual(paragraph_text(table["rows"][1][0]["blocks"][0]), "대니")
        self.assertEqual(table["rows"][0][1]["blocks"][0]["alignment"], "right")
        self.assertTrue(table["rows"][0][0]["blocks"][0]["runs"][0]["bold"])
        self.assertEqual(parsed["warnings"], [])
        json.dumps(parsed, ensure_ascii=False)

    def test_markdown_image_order_remote_warning_and_page_break(self):
        self.image()
        parsed = self.markdown(
            "앞 ![설명](%EA%B7%B8%EB%A6%BC.png) 뒤\n\n"
            "![원격](https://example.invalid/image.png)\n\n"
            "<!-- pagebreak -->\n\n다음 페이지\n")
        blocks = parsed["blocks"]
        self.assertEqual([b["type"] for b in blocks[:3]], ["paragraph", "image", "paragraph"])
        self.assertEqual(paragraph_text(blocks[0]), "앞 ")
        self.assertEqual(paragraph_text(blocks[2]), " 뒤")
        image = blocks[1]
        self.assertEqual(image["alt"], "설명")
        self.assertEqual((image["width_pt"], image["height_pt"]), (60, 30))
        self.assertTrue(Path(image["path"]).is_file())
        self.assertTrue(Path(image["path"]).is_relative_to(self.assets.resolve()))
        self.assertTrue(any("not fetched" in w for w in parsed["warnings"]))
        self.assertTrue(any(b.get("page_break_before") for b in blocks))

    def test_missing_image_task_list_and_html_are_explicit(self):
        parsed = self.markdown("- [x] 끝\n- [ ] 대기\n\n![없음](missing.png)\n\n<div>내용</div>\n")
        self.assertEqual([b.get("list_prefix") for b in parsed["blocks"][:2]], ["☑", "☐"])
        self.assertTrue(any("could not be read" in w for w in parsed["warnings"]))
        self.assertTrue(any("HTML" in w for w in parsed["warnings"]))
        self.assertTrue(any("없음" in paragraph_text(b) for b in parsed["blocks"]
                            if b["type"] == "paragraph"))

    def test_docx_block_order_styles_and_embedded_image(self):
        document = Document()
        document.add_heading("한글 제목", 2)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("포맷")
        run.bold = True
        run.italic = True
        run.underline = True
        run.font.strike = True
        run.font.size = Pt(15)
        run.font.name = "Arial"
        run.font.color.rgb = RGBColor(12, 34, 56)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "표 내용"
        table.cell(0, 1).text = "42"
        paragraph = document.add_paragraph("앞")
        paragraph.add_run().add_picture(str(self.image()), width=Inches(1))
        paragraph.add_run("뒤")
        document.add_paragraph("끝")
        parsed = self.save_docx(document)
        blocks = parsed["blocks"]
        self.assertEqual([b["type"] for b in blocks],
                         ["paragraph", "paragraph", "table", "paragraph", "image", "paragraph", "paragraph"])
        self.assertEqual(blocks[0]["heading"], 2)
        self.assertEqual(paragraph_text(blocks[0]), "한글 제목")
        formatted = blocks[1]["runs"][0]
        self.assertTrue(all(formatted[k] for k in ("bold", "italic", "underline", "strike")))
        self.assertEqual((formatted["size"], formatted["font"], formatted["color"]),
                         (15, "Arial", "0C2238"))
        self.assertEqual(blocks[1]["alignment"], "center")
        self.assertEqual((blocks[4]["width_pt"], blocks[4]["height_pt"]), (72, 36))
        self.assertTrue(Path(blocks[4]["path"]).is_file())
        self.assertEqual(paragraph_text(blocks[3]), "앞")
        self.assertEqual(paragraph_text(blocks[5]), "뒤")
        self.assertEqual(parsed["warnings"], [])

    def test_docx_merged_table_avoids_duplicate_text(self):
        document = Document()
        table = document.add_table(rows=3, cols=3)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "가로"
        table.cell(1, 0).merge(table.cell(2, 0)).text = "세로"
        table.cell(1, 1).text = "옆"
        table.cell(2, 1).text = "아래"
        result = self.save_docx(document)["blocks"][0]
        self.assertEqual(result["rows"][0][0]["colspan"], 2)
        self.assertEqual(result["rows"][1][0]["rowspan"], 2)
        text = "|".join(paragraph_text(p) for row in result["rows"] for cell in row
                        for p in cell["blocks"])
        self.assertEqual(text.count("세로"), 1)
        self.assertEqual(len(result["rows"][2]), 2)

    def test_docx_page_breaks_and_list_style(self):
        document = Document()
        document.add_paragraph("목록", style="List Number")
        document.add_paragraph("두번째", style="List Number")
        paragraph = document.add_paragraph("이전")
        paragraph.add_run().add_break(WD_BREAK.PAGE)
        paragraph.add_run("다음")
        paragraph = document.add_paragraph("새 페이지")
        paragraph.paragraph_format.page_break_before = True
        parsed = self.save_docx(document)
        paragraphs = parsed["blocks"]
        self.assertEqual([p.get("list_prefix") for p in paragraphs[:2]], ["1.", "2."])
        self.assertEqual(paragraph_text(paragraphs[2]), "이전")
        self.assertEqual(paragraph_text(paragraphs[3]), "다음")
        self.assertTrue(paragraphs[3]["page_break_before"])
        self.assertTrue(paragraphs[4]["page_break_before"])

    def test_docx_unsupported_features_are_warned_and_hyperlink_text_retained(self):
        document = Document()
        document.sections[0].header.paragraphs[0].text = "머리말"
        paragraph = document.add_paragraph("본문 ")
        hyperlink = OxmlElement("w:hyperlink")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "링크"
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        footnote = OxmlElement("w:footnoteReference")
        footnote.set(qn("w:id"), "1")
        paragraph.add_run()._r.append(footnote)
        textbox = OxmlElement("w:pict")
        paragraph.add_run()._r.append(textbox)
        parsed = self.save_docx(document)
        self.assertEqual(paragraph_text(parsed["blocks"][0]), "본문 링크")
        warnings = parsed["warnings"]
        self.assertEqual(paragraph_text(parsed["header"][0]), "머리말")
        self.assertTrue(any("unresolved DOCX links" in w for w in warnings))
        self.assertTrue(any("unresolved DOCX footnote" in w for w in warnings))
        self.assertTrue(any("pict" in w for w in warnings))

    def test_image_asset_symlink_rejected(self):
        self.image()
        parsed = self.markdown("![그림](그림.png)")
        target = Path(parsed["blocks"][0]["path"])
        target.unlink()
        other = self.directory / "outside.txt"
        other.write_text("unchanged")
        target.symlink_to(other)
        with self.assertRaisesRegex(ValueError, "symlink"):
            self.markdown("![그림](그림.png)")
        self.assertEqual(other.read_text(), "unchanged")

    def test_unknown_extension_and_missing_file_fail(self):
        unknown = self.directory / "input.txt"
        unknown.write_text("x")
        with self.assertRaises(ValueError):
            parse_input(unknown, self.assets)
        with self.assertRaises(FileNotFoundError):
            parse_input(self.directory / "missing.md", self.assets)


if __name__ == "__main__":
    unittest.main()
