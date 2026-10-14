"""Native HWP round-trip tests for report layout, controls and templates."""
import subprocess
import sys
import tempfile
import unittest
from lxml import etree
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / "vendor/python")]
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from convert import convert, inspect_hwp, java_home


def controls(layout):
    for p in layout:
        for c in p["controls"]:
            yield c
            yield from controls(c.get("paragraphs", []))
            for cell in c.get("cells", []):
                yield from controls(cell["paragraphs"])


class FidelityTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def layout_doc(self):
        d = Document()
        d.sections[0].left_margin = Inches(1.25)
        d.sections[0].right_margin = Inches(.5)
        p = d.add_paragraph("본문 ")
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(18)
        p.paragraph_format.line_spacing = 1.25
        p.add_run("3").font.superscript = True
        t = d.add_table(rows=1, cols=2)
        t.columns[0].width = Inches(1)
        t.columns[1].width = Inches(4)
        t.cell(0, 0).text, t.cell(0, 1).text = "번호", "긴 설명"
        return d

    def save(self, d, name="input.docx"):
        p = self.root / name
        d.save(p)
        return p

    def test_layout_and_unequal_widths_survive_native_roundtrip(self):
        source = self.save(self.layout_doc())
        h = inspect_hwp(convert(source)["output"])
        self.assertEqual((h["page"]["left"], h["page"]["right"]), (9000, 3600))
        self.assertEqual((h["layout"][0]["before"], h["layout"][0]["after"]), (700, 1800))
        self.assertEqual(h["layout"][0]["line_spacing"], 125)
        self.assertTrue(any(c["super"] for c in h["layout"][0]["characters"]))
        table = next(c for c in controls(h["layout"]) if c["type"] == "ControlTable")
        a, b = (c["width"] for c in table["cells"])
        self.assertLessEqual(abs(b - 4*a), 4)
        self.assertEqual((a,b), (7200,28800))
        self.assertEqual(h["page"]["top"]+h["page"]["header"], 7200)
        self.assertEqual(h["page"]["bottom"]+h["page"]["footer"], 7200)

    def test_footnote_hyperlink_and_page_field_are_native_controls(self):
        d = Document(); p = d.add_paragraph("본문")
        note = OxmlElement("w:footnoteReference"); note.set(qn("w:id"), "1")
        p.add_run()._r.append(note)
        notes = OxmlElement("w:footnotes")
        fn = OxmlElement("w:footnote"); fn.set(qn("w:id"), "1")
        np = OxmlElement("w:p"); nr = OxmlElement("w:r"); nt = OxmlElement("w:t")
        nt.text = "각주 내용 보존"; nr.append(nt); np.append(nr); fn.append(np); notes.append(fn)
        part = Part(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, etree.tostring(notes), d.part.package)
        d.part.relate_to(part, RT.FOOTNOTES)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), d.part.relate_to("https://example.com/a?b=1", RT.HYPERLINK, is_external=True))
        lr = OxmlElement("w:r"); lt = OxmlElement("w:t"); lt.text = "출처 링크"; lr.append(lt); link.append(lr); p._p.append(link)
        footer = d.sections[0].footer.paragraphs[0]
        field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), " PAGE ")
        r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "999"; r.append(t); field.append(r); footer._p.append(field)
        report = convert(self.save(d))
        self.assertEqual((report["footnotes"], report["hyperlinks"], report["page_fields"]), (1, 1, 1))
        cs = list(controls(inspect_hwp(report["output"])["layout"]))
        fn = next(c for c in cs if c["type"] == "ControlFootnote")
        self.assertIn("각주 내용 보존", fn["paragraphs"][0]["text"])
        self.assertTrue(any(c.get("number_sort") == "Page" for c in cs))
        self.assertTrue(any("example.com" in c.get("command", "") for c in cs))
        footer = next(c for c in cs if c["type"] == "ControlFooter")
        self.assertNotIn("999", str(footer))

    def test_reference_style_and_page_override_without_copying_body(self):
        source = self.save(self.layout_doc())
        seed = self.root / "seed.md"; seed.write_text("템플릿 기존본문")
        seed_hwp = convert(seed)["output"]
        reference = self.root / "reference.hwp"
        subprocess.run([str(java_home()/"bin/java"), "--class-path", str(ROOT/"vendor/*"),
                        str(ROOT/"tests/TemplateFixture.java"), seed_hwp, str(reference)], check=True, capture_output=True)
        output = self.root / "templated.hwp"
        convert(source, output, reference_doc=reference)
        h = inspect_hwp(output)
        self.assertEqual(h["page"]["left"], 9000)
        self.assertEqual(h["layout"][0]["after"], 2400)
        self.assertTrue(all(c["size"] == 1800 for c in h["layout"][0]["characters"]))
        self.assertNotIn("템플릿 기존본문", h["text"])
        self.assertIn("긴 설명", h["text"])

    def test_image_template_is_rejected_without_overwriting_output(self):
        from PIL import Image
        image = self.root / "image.png"
        Image.new("RGB", (8, 8), "blue").save(image)
        seed = self.root / "seed.md"
        seed.write_text("![logo](image.png)")
        reference = convert(seed)["output"]
        output = self.root / "protected.hwp"
        output.write_bytes(b"existing document")
        with self.assertRaises(Exception):
            convert(self.save(Document()), output, reference_doc=reference, force=True)
        self.assertEqual(output.read_bytes(), b"existing document")

    def test_bad_reference_does_not_publish_output(self):
        source = self.save(Document())
        output = self.root / "bad.hwp"
        with self.assertRaises(ValueError):
            convert(source, output, reference_doc=self.root/"missing.hwp")
        self.assertFalse(output.exists())


if __name__ == "__main__":
    unittest.main()
